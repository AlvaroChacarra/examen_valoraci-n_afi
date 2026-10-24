#!/usr/bin/env python3
"""Independent source/Markdown conformance checks for the second review.

No prices are computed and no workbook is modified or recalculated.
"""
import argparse,collections,hashlib,io,json,re,warnings
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as E
from urllib.parse import unquote
import openpyxl
from normalize_sources import ROOT,defined_names_xml,formula_text

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--archives-dir',type=Path);args=parser.parse_args()
    warnings.simplefilter('ignore');sources=json.loads((ROOT/'catalogo.json').read_text())['sources']
    stats=collections.Counter();errors=[];graph=collections.defaultdict(set)
    for r in sources:
        path=ROOT/r['path'];data=path.read_bytes();assert hashlib.sha256(data).hexdigest()==r['sha256'],r['id']
        text=(ROOT/r['extraction']).read_text();stats['sources']+=1
        if r['extension'] not in ['.xlsx','.xlsm']:continue
        z=ZipFile(path);xml=E.fromstring(z.read('xl/workbook.xml'));sheets=xml.findall('.//{*}sheets/{*}sheet')
        sections=re.split(r'^## Hoja: (.+)$',text,flags=re.M);byname={sections[i]:sections[i+1] for i in range(1,len(sections),2)}
        wb=openpyxl.load_workbook(path,data_only=False,read_only=False,keep_links=True)
        xml_count=sum(len(E.fromstring(z.read(n)).findall('.//{*}sheetData/{*}row/{*}c/{*}f')) for n in z.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n))
        assert xml_count==r['formulas'],(r['id'],'raw formula count')
        for s in sheets:
            name=s.get('name');state=s.get('state','visible');body=byname[name];stats['sheets']+=1
            assert 'Estado: `'+state+'`' in body,(r['id'],name,'state')
            stats['hidden_sheets']+=state!='visible'
            rows=re.findall(r'^\| `([A-Z]+\d+)` \| `([A-Z]+\d+)` \| (\d+) \| `(.*?)` \|',body,re.M)
            ws=wb[name];fc=sum(c.data_type=='f' for c in ws._cells.values())
            assert sum(int(a[2]) for a in rows)==fc,(r['id'],name,'family coverage')
            for first,last,count,formula in rows:
                actual=formula_text(ws[first].value).replace('|','\\|').replace('\n','<br>')
                assert actual==formula,(r['id'],name,first,'formula fidelity')
                assert ws[last].data_type=='f',(r['id'],name,last,'last formula absent')
                stats['formula_families']+=1
            stats['formulas']+=fc
        names=defined_names_xml(path)
        head=text.split('## Hoja:',1)[0]
        for n in names:
            needle=f'| `{n["name"]}` | `{n["scope"]}` | `{n["ref"].replace(chr(124),chr(92)+chr(124))}` |'
            # Books without locals retain the original two-column name table.
            if not any(x['local_sheet_id'] is not None for x in names):needle=f'| `{n["name"]}` | `{n["ref"]}` |'
            assert needle in head,(r['id'],n['name'],n['scope'],'name omitted')
            stats['defined_names']+=1;stats['local_names']+=n['local_sheet_id'] is not None
        wb.close()
    for sid in ['PDF-005','PDF-006']:
        r=next(x for x in sources if x['id']==sid);text=(ROOT/r['extraction']).read_text()
        assert text.count('### OCR auxiliar')==r['pages']==r['ocr_pages'],sid
        assert len(re.findall(r'^## Página \d+$',text,re.M))==r['pages'],sid
        stats['ocr_pages']+=r['pages']
    for path in ROOT.rglob('*.md'):
        if any(x in path.relative_to(ROOT).parts for x in ['metadata','.git']):continue
        text=re.sub(r'```.*?```','',path.read_text(),flags=re.S)
        for target in re.findall(r'!?\[[^\]]*\]\(([^)\n]+)\)',text):
            if re.match(r'^(https?:|mailto:|#)',target):continue
            target=target.split('#',1)[0]
            if not target:continue
            resolved=(path.parent/unquote(target)).resolve();assert resolved.exists(),(path,target)
            stats['links']+=1;graph[path.resolve()].add(resolved)
    visited=set();pending=[(ROOT/'AGENTS.md').resolve()]
    while pending:
        p=pending.pop()
        if p in visited:continue
        visited.add(p);pending.extend(graph[p]-visited)
    orphans=[p for p in (ROOT/'knowledge-base').rglob('*.md') if p.resolve() not in visited]
    assert not orphans,orphans
    if args.archives_dir:
        received=[]
        def walk(data):
            for item in ZipFile(io.BytesIO(data)).infolist():
                if item.is_dir() or '__MACOSX' in item.filename:continue
                with ZipFile(io.BytesIO(data)) as z:content=z.read(item)
                if item.filename.lower().endswith('.zip'):walk(content)
                else:received.append(hashlib.sha256(content).hexdigest())
        for p in args.archives_dir.glob('*.zip'):walk(p.read_bytes())
        expected=[r['sha256'] for r in sources if r['group']!='encargo']
        assert collections.Counter(received)==collections.Counter(expected),'archive leaf coverage'
        stats['archive_documents']=len(received)
    assert stats['sources']==59 and stats['sheets']==60 and stats['formulas']==60701
    print(json.dumps({'status':'PASS','checks':dict(stats),'broken_links':0,'orphan_markdown':0,'scope':'documentary readiness; no repricing'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
