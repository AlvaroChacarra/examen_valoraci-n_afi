"""Recálculo real con LibreOffice y cotejo de resultados/estados con Python."""
from pathlib import Path
import argparse,json,subprocess,shutil,time
import openpyxl
import numpy as np
from .ejercicio1 import Tree,EXERCISE
ROOT=Path(__file__).resolve().parents[1]

def verify(folder):
 out={}
 for filename in ['Ejercicio1.xlsx','Ejercicio2.xlsx','Ejercicio3.xlsx']:
  book=openpyxl.load_workbook(folder/filename,data_only=True,read_only=True)
  errors=[];cells=0
  for sheet in book:
   for row in sheet:
    for cell in row:
     if cell.data_type=='e':errors.append(f'{sheet.title}!{cell.coordinate}: {cell.value}')
     cells+=cell.value is not None
  assert not errors, errors[:20]
  out[filename]={'nonempty_cells':cells,'formula_errors':errors}
  if filename=='Ejercicio1.xlsx':
   maxdiff=0.;nodecount=0
   for suffix,elapsed,spot in [('0',0.,50.),('15',15/365,52.5)]:
    trees=[Tree(spot=spot,elapsed=elapsed,style=sty) for sty in ['european','bermuda','american']];smooth=Tree(spot=spot,elapsed=elapsed,smooth=True)
    rows=iter(book[f'Arbol{suffix}'].iter_rows(min_row=2,values_only=True))
    for j,l in enumerate(trees[1].layers):
     for ix in np.ndindex(l.spot.shape):
      row=next(rows);expected=[l.spot[ix],*(t.layers[j].value[ix] for t in trees),smooth.layers[j].value[ix]];actual=[row[5],*row[9:13]]
      diff=max(abs(float(a)-float(b)) for a,b in zip(actual,expected));maxdiff=max(maxdiff,diff);nodecount+=1
      assert bool(row[16])==bool(l.stop[ix]),(suffix,j,ix,'politica')
      if j<36:
       assert abs(row[18])<1e-10 and abs(row[19])<1e-10
    assert maxdiff<1e-9,maxdiff
   exact=Tree().exact_exercise()
   for row,j in zip(book['Probabilidades'].iter_rows(min_row=2,max_row=15,values_only=True),sorted(EXERCISE)):
    r=exact['by_month'][j]
    assert abs(row[1]-r['survival'])<1e-12 and abs(row[2]-r['first'])<1e-12 and abs(row[3]-r['conditional'])<1e-12
   out[filename].update(max_node_error=maxdiff,states_verified=nodecount,probabilities_verified=14)
  else:
   e=int(filename[9]);ref=json.loads((ROOT/f'resultados/ejercicio{e}.json').read_text());s=book['Calculo'];pairs=[(17,'fra'),(18,'caplet'),(19,'floorlet')] if e==2 else [(24,'forward'),(25,'forward_quanto'),(26,'call'),(27,'call_quanto'),(28,'put'),(29,'put_quanto')]
   diffs={key:abs(s[f'B{row}'].value-ref['closed'][key]) for row,key in pairs}
   assert max(diffs.values())<1e-7,diffs
   out[filename]['price_errors']=diffs
  book.close()
 return out

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--recalculate',action='store_true');ap.add_argument('--folder',type=Path,default=ROOT/'hojas');args=ap.parse_args()
 if args.recalculate:
  dest=ROOT.parent/'tmp/recalculadas';dest.mkdir(parents=True,exist_ok=True)
  for i in [1,2,3]:
   target=dest/f'Ejercicio{i}.xlsx'
   if target.exists():target.unlink()
  cmd=['soffice','-env:UserInstallation=file:///tmp/afi-lo-repro','--headless','--convert-to','xlsx','--outdir',str(dest),*[str(ROOT/'hojas'/f'Ejercicio{i}.xlsx') for i in [1,2,3]]]
  subprocess.run(cmd,check=True,timeout=300)
  # Detectar fallos aunque LibreOffice retorne cero.
  for i in [1,2,3]:assert (dest/f'Ejercicio{i}.xlsx').stat().st_size>0
  out=verify(dest)
  for i in [1,2,3]:shutil.copy2(dest/f'Ejercicio{i}.xlsx',ROOT/'hojas')
 else:out=verify(args.folder)
 (ROOT/'resultados/hojas_verificadas.json').write_text(json.dumps(out,indent=2));print(json.dumps(out,indent=2))
if __name__=='__main__':main()
